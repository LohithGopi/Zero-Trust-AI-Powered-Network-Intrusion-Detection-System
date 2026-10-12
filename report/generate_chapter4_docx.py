import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

REPORT_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(REPORT_DIR, "Chapter_4_Results_and_Analysis.md")
FIG_DIR = os.path.join(REPORT_DIR, "figures")
OUT_DOCX = os.path.join(REPORT_DIR, "Chapter_4_Results_and_Analysis.docx")


def clean_inline_for_docx(text: str) -> str:
    """Clean LaTeX math and Markdown tokens into readable Unicode for Word."""
    s = text.strip()
    replacements = [
        (r"$\|\Delta W^{(e)}\|_F = \sum_l \|W^{(l,e)} - W^{(l,e-1)}\|_F$", "||ΔW^(e)||_F = Σ_l ||W^(l,e) - W^(l,e-1)||_F"),
        (r"$\|\Delta W^{(e)}\|_F$", "||ΔW^(e)||_F"),
        (r"$\|\Delta W^{(1)}\|_F$", "||ΔW^(1)||_F"),
        (r"$\|\Delta W\|_F$", "||ΔW||_F"),
        (r"$\Delta W^{(e)}$", "ΔW^(e)"),
        (r"$\Delta W$", "ΔW"),
        (r"$\Delta t = 0$", "Δt = 0"),
        (r"$N \times 1 \times 122$", "N × 1 × 122"),
        (r"$X_{train}, y_{train}$", "X_train, y_train"),
        (r"$X_{val}, y_{val}$", "X_val, y_val"),
        (r"$X_{test}, y_{test}$", "X_test, y_test"),
        (r"$X_{train\_raw}$", "X_train_raw"),
        (r"$X_{val\_raw}$", "X_val_raw"),
        (r"$X_{test\_raw}$", "X_test_raw"),
        (r"$X_{train}, X_{val}, X_{test}$", "X_train, X_val, X_test"),
        (r"$\mu_{train}, \sigma_{train}$", "μ_train, σ_train"),
        (r"$\mu \approx 0, \sigma \approx 1$", "μ ≈ 0, σ ≈ 1"),
        (r"$W^{(1)} \in \mathbb{R}^{d \times 64}$", "W^(1) ∈ R^(d×64)"),
        (r"$W^{(2)} \in \mathbb{R}^{64 \times 32}$", "W^(2) ∈ R^(64×32)"),
        (r"$W^{(3)} \in \mathbb{R}^{32 \times C}$", "W^(3) ∈ R^(32×C)"),
        (r"$d \in \{15, 19, 22\}$", "d ∈ {15, 19, 22}"),
        (r"$X_{raw} \in \mathbb{R}^{1 \times d}$", "X_raw ∈ R^(1×d)"),
        (r"$Z = (X_{raw} - \mu_{train}) / \sigma_{train}$", "Z = (X_raw - μ_train) / σ_train"),
        (r"$\hat{p} \in [0, 1]^C$", "p̂ ∈ [0, 1]^C"),
        (r"$\hat{y} = \arg\max_c \hat{p}_c$", "ŷ = argmax_c p̂_c"),
        (r"$\hat{L}$", "L̂"),
        (r"$\text{Conf} < 0.60$", "Conf < 0.60"),
        (r"$0.60 \le \text{Conf} < 0.85$", "0.60 ≤ Conf < 0.85"),
        (r"$\text{Conf} \ge 0.85$", "Conf ≥ 0.85"),
        (r"$ROC\text{-}AUC$", "ROC-AUC"),
        (r"$F_1\text{-Score}$", "F1-Score"),
        (r"$F_1$", "F1"),
        (r"$FPR < 2.5\%$", "FPR < 2.5%"),
        (r"$FPR$", "FPR"),
        (r"$FNR$", "FNR"),
        (r"$TPR$", "TPR"),
        (r"$TP + FP = 0$", "TP + FP = 0"),
        (r"$FN = 930, TP = 0$", "FN = 930, TP = 0"),
        (r"$TP$", "TP"),
        (r"$TN$", "TN"),
        (r"$FP$", "FP"),
        (r"$FN$", "FN"),
        (r"$N_{test} = 3,750$", "N_test = 3,750"),
        (r"$N_{test}$", "N_test"),
        (r"$\ge 60\%$", "≥ 60%"),
        (r"$\ge 85\%$", "≥ 85%"),
        (r"$\ge$", "≥"),
        (r"$\le$", "≤"),
        (r"$\rightarrow$", "→"),
        (r"$\times$", "×"),
    ]
    for old, new in replacements:
        s = s.replace(old, new)

    s = re.sub(r"\$([^$]+)\$", r"\1", s)
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"\*(.*?)\*", r"\1", s)
    s = re.sub(r"`(.*?)`", r"\1", s)
    return s


def build_docx_from_markdown():
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Zero Trust AI-Powered Network Intrusion Detection System — Chapter 4: Results and Analysis"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Times New Roman"
        hp.style.font.size = Pt(9)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Department of Information Science and Engineering, JNNCE Shivamogga"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style.font.name = "Times New Roman"
        fp.style.font.size = Pt(9)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11.5)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(5)

    with open(MD_PATH, "r", encoding="utf-8") as f:
        lines = f.readlines()

    table_rows = []
    quote_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        clean_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip().strip("|").split("|")]
            if all(set(c) <= set(":- ") for c in cells):
                continue
            clean_rows.append(cells)
        if clean_rows:
            cols = max(len(r) for r in clean_rows)
            tbl = doc.add_table(rows=len(clean_rows), cols=cols)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(clean_rows):
                for c_idx in range(min(cols, len(row_data))):
                    cell = tbl.cell(r_idx, c_idx)
                    txt = clean_inline_for_docx(row_data[c_idx])
                    p = cell.paragraphs[0]
                    run = p.add_run(txt)
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
                    if r_idx == 0:
                        run.bold = True
            doc.add_paragraph()
        table_rows = []

    def flush_quote():
        nonlocal quote_rows
        if not quote_rows:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        first_p = cell.paragraphs[0]
        for idx, q in enumerate(quote_rows):
            txt = clean_inline_for_docx(q.lstrip(">").strip())
            p = first_p if idx == 0 else cell.add_paragraph()
            run = p.add_run(txt)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.0)
            if idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(29, 78, 216)
        doc.add_paragraph()
        quote_rows = []

    for raw_line in lines:
        line = raw_line.rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            if quote_rows:
                flush_quote()
            table_rows.append(stripped)
            continue
        elif table_rows:
            flush_table()

        if stripped.startswith(">"):
            quote_rows.append(stripped)
            continue
        elif quote_rows:
            flush_quote()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("$$") and stripped.endswith("$$"):
            raw = stripped.strip("$").strip()
            if "Accuracy" in raw:
                eq_txt = "Accuracy = (TP + TN) / (TP + TN + FP + FN)   |   Precision = TP / (TP + FP)   |   Recall = TP / (TP + FN)"
            elif "F_1" in raw:
                eq_txt = "F1-Score = (2 × Precision × Recall) / (Precision + Recall)   |   FPR = FP / (FP + TN)"
            else:
                eq_txt = clean_inline_for_docx(raw)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(eq_txt)
            run.font.name = "Cambria Math"
            run.font.size = Pt(10.5)
            run.bold = True
            continue

        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if img_match:
            caption_txt = clean_inline_for_docx(img_match.group(1))
            img_rel = img_match.group(2)
            img_path = os.path.join(REPORT_DIR, img_rel.replace("/", os.sep))
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(6.2))
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                crun = cp.add_run(caption_txt)
                crun.font.name = "Times New Roman"
                crun.font.size = Pt(10)
                crun.bold = True
                crun.italic = True
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(clean_inline_for_docx(stripped[2:]), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            h = doc.add_heading(clean_inline_for_docx(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            h = doc.add_heading(clean_inline_for_docx(stripped[4:]), level=2)
        elif stripped.startswith("**Table 4."):
            p = doc.add_paragraph()
            run = p.add_run(clean_inline_for_docx(stripped))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        elif stripped.startswith("- ") or re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(clean_inline_for_docx(stripped))
            p.paragraph_format.left_indent = Inches(0.25)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p = doc.add_paragraph(clean_inline_for_docx(stripped))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    flush_table()
    flush_quote()

    doc.save(OUT_DOCX)
    print(f"Successfully generated Word document: {OUT_DOCX}")


if __name__ == "__main__":
    build_docx_from_markdown()
