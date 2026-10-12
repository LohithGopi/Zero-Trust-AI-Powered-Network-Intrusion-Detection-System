import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

REPORT_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(REPORT_DIR, "Chapter_3_System_Design_and_Implementation.md")
FIG_DIR = os.path.join(REPORT_DIR, "figures")
OUT_DOCX = os.path.join(REPORT_DIR, "Chapter_3_System_Design_and_Implementation.docx")


def build_docx_from_markdown():
    doc = Document()

    # Configure standard A4 page margins (1 inch)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "AI-Powered Zero Trust Network Intrusion Detection System — Chapter 3: System Design and Implementation"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Times New Roman"
        hp.style.font.size = Pt(9)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Department of Information Science and Engineering, JNNCE Shivamogga (Batch No. 34)"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style.font.name = "Times New Roman"
        fp.style.font.size = Pt(9)

    # Base Normal style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    with open(MD_PATH, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    code_buffer = []
    in_table = False
    table_rows = []
    in_math = False
    math_buffer = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        # Filter out separator row | :--- |
        clean_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip().strip("|").split("|")]
            if all(set(c) <= set(":- ") for c in cells):
                continue
            clean_rows.append(cells)
        if clean_rows:
            cols = max(len(r) for r in clean_rows)
            tbl = doc.add_table(rows=len(clean_rows), cols=cols)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(clean_rows):
                for c_idx in range(min(cols, len(row_data))):
                    cell = tbl.cell(r_idx, c_idx)
                    txt = re.sub(r"\*\*(.*?)\*\*", r"\1", row_data[c_idx])
                    txt = re.sub(r"`(.*?)`", r"\1", txt)
                    p = cell.paragraphs[0]
                    run = p.add_run(txt)
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)
                    if r_idx == 0:
                        run.bold = True
            doc.add_paragraph()
        table_rows = []
        in_table = False

    for raw_line in lines:
        line = raw_line.rstrip("\n")

        # Handle fenced code blocks
        if line.strip().startswith("```"):
            if in_table:
                flush_table()
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_buffer))
                run.font.name = "Courier New"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
            continue

        if in_code:
            code_buffer.append(line)
            continue

        # Handle display math blocks $$ ... $$
        if line.strip().startswith("$$"):
            if in_table:
                flush_table()
            if line.strip().endswith("$$") and len(line.strip()) > 4:
                eq_txt = line.strip()[2:-2].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(eq_txt)
                run.font.name = "Cambria Math"
                run.font.size = Pt(11)
                run.bold = True
                continue
            else:
                if not in_math:
                    in_math = True
                    math_buffer = [line.strip().lstrip("$")]
                else:
                    in_math = False
                    math_buffer.append(line.strip().rstrip("$"))
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(" ".join(math_buffer))
                    run.font.name = "Cambria Math"
                    run.font.size = Pt(11)
                    run.bold = True
                continue

        if in_math:
            math_buffer.append(line.strip())
            continue

        # Handle Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_rows.append(line)
            continue
        elif in_table:
            flush_table()

        if not line.strip() or line.strip() == "---":
            continue

        # Handle images ![alt](figures/...)
        img_match = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if img_match:
            img_rel = img_match.group(2)
            img_path = os.path.join(REPORT_DIR, img_rel.replace("/", os.sep))
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(6.1))
            continue

        # Headings
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_page_break()
            h = doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            h = doc.add_heading(line[5:].strip(), level=3)
        elif line.strip().startswith("*Figure "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip().strip("*"))
            run.italic = True
            run.bold = True
            run.font.size = Pt(10.5)
        elif line.strip().startswith("- ") or re.match(r"^\d+\.\s", line.strip()):
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", line.strip())
            clean = re.sub(r"`(.*?)`", r"\1", clean)
            p = doc.add_paragraph(clean)
            p.paragraph_format.left_indent = Inches(0.25)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", line.strip())
            clean = re.sub(r"`(.*?)`", r"\1", clean)
            p = doc.add_paragraph(clean)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if in_table:
        flush_table()

    doc.save(OUT_DOCX)
    print(f"Successfully generated Word document: {OUT_DOCX}")


if __name__ == "__main__":
    build_docx_from_markdown()
